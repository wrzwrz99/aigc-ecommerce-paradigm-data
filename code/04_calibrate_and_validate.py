# -*- coding: utf-8 -*-
"""校准定稿：重算 150篇指标 + 全样本分布/年度/描述统计 -> 存档"""
import pandas as pd
from collections import Counter, defaultdict
import sys
sys.path.insert(0, r'D:\论文\电商综述\AIGC文创电商\查新文本')
import _train_worddict2 as tw

CLASSES = tw.CLASSES

def kappa(y1, y2):
    n = len(y1)
    p_o = sum(a == b for a, b in zip(y1, y2)) / n
    cats = set(y1) | set(y2)
    p_e = 0
    for c in cats:
        p_e += (sum(1 for x in y1 if x == c) / n) * (sum(1 for x in y2 if x == c) / n)
    return (p_o - p_e) / (1 - p_e) if p_e != 1 else 1

cr = pd.read_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_clean_records.csv', encoding='utf-8-sig')
mp = pd.read_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_dict_codes_for_compare.csv', encoding='utf-8-sig')
import openpyxl
wb = openpyxl.load_workbook(r'D:\论文\电商综述\AIGC文创电商\查新文本\人工复核表_字典编码_150篇.xlsx')
ws = wb['人工复核表']
manual = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=5).value for r in range(3, 153)}
mp['manual'] = mp['序号'].map(manual)
tr = cr.merge(mp[['题名', 'manual']], on='题名', how='inner')

# 150篇评估
preds = [tw.assign(tw.score_text(r['摘要'], r['题名'])) for _, r in tr.iterrows()]
acc = sum(a == b for a, b in zip(preds, tr['manual'])) / len(tr)
kap = kappa(preds, tr['manual'])
cm = defaultdict(Counter)
for t, p in zip(tr['manual'], preds):
    cm[t][p] += 1
f1s = {}
for c in CLASSES:
    tp = cm[c][c]
    prec = tp / sum(cm[x][c] for x in CLASSES) if sum(cm[x][c] for x in CLASSES) else 0
    rec = tp / sum(cm[c][x] for x in CLASSES) if sum(cm[c][x] for x in CLASSES) else 0
    f1s[c] = 2 * prec * rec / (prec + rec) if prec + rec else 0

# 全样本预测
allp = [tw.assign(tw.score_text(r['摘要'], r['题名'])) for _, r in cr.iterrows()]
pc = Counter(allp)
cr2 = cr.copy()
cr2['范式_校准'] = allp

# 年度 x 范式
yr = cr2.pivot_table(index='年份', columns='范式_校准', values='题名', aggfunc='count', fill_value=0)
yr['合计'] = yr.sum(axis=1)

# 描述统计：范式 x 引用/IF（若数据存在）
nb = pd.read_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_nb_data_clean.csv', encoding='utf-8-sig')
nb2 = nb.merge(cr2[['题名', '范式_校准']], on='题名', how='left')
desc = nb2.groupby('范式_校准')[['ts_cit', 'cite_cnt', 'if_val']].mean() if 'if_val' in nb2.columns else None

with open(r'C:\Temp_rend\_cv_result.txt', 'w', encoding='utf-8') as f:
    f.write('===== 最终校准（词级词典 v3 + 非逻辑须>逻辑规则） =====\n')
    f.write('150篇人工复核: 准确率=%.3f Cohen kappa=%.3f\n' % (acc, kap))
    f.write('混淆矩阵(行=人工 列=词典):\n')
    f.write('  ' + ''.join('%10s' % c for c in CLASSES) + '\n')
    for c in CLASSES:
        f.write('  %-4s' % c + ''.join('%10d' % cm[c][x] for x in CLASSES) + '\n')
    for c in CLASSES:
        tp = cm[c][c]
        prec = tp / sum(cm[x][c] for x in CLASSES) if sum(cm[x][c] for x in CLASSES) else 0
        rec = tp / sum(cm[c][x] for x in CLASSES) if sum(cm[c][x] for x in CLASSES) else 0
        f.write('  %s: P=%.3f R=%.3f F1=%.3f (tp=%d)\n' % (c, prec, rec, f1s[c], tp))
    f.write('全样本分布: %s\n' % dict(pc))
    for c in CLASSES:
        f.write('  %s: %d (%.1f%%)\n' % (c, pc[c], 100 * pc[c] / len(allp)))
    f.write('\n年度x范式:\n')
    f.write(yr.to_string() + '\n')
    if desc is not None:
        f.write('\n范式x引用/IF 均值:\n')
        f.write(desc.to_string() + '\n')

# 存档关键文件
cr2[['题名', '期刊', '年份', '范式_校准', '摘要']].to_csv(
    r'D:\论文\电商综述\AIGC文创电商\查新文本\_calibrated_codes.csv', index=False, encoding='utf-8-sig')
# 150篇比对表
tr2 = mp[['序号', '题名']].copy()
tr2 = tr2.merge(tr[['题名', 'manual']], on='题名', how='left')
tr2['人工编码'] = list(tr2['manual'])
tr2['词典编码'] = preds
tr2['一致'] = [a == b for a, b in zip(tr2['人工编码'], tr2['词典编码'])]
tr2.to_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_calibration_150_compare.csv',
           index=False, encoding='utf-8-sig')
yr.to_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_year_paradigm_calibrated.csv', encoding='utf-8-sig')

print('准确率=%.3f kappa=%.3f' % (acc, kap))
print('F1:', {c: round(f1s[c], 3) for c in CLASSES})
print('全样本分布:', dict(pc))
print('年度表:')
print(yr)
if desc is not None:
    print('描述统计:')
    print(desc)
print('已存档')
