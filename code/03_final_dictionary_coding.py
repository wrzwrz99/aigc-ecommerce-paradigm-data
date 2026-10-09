# -*- coding: utf-8 -*-
"""词级词典 v2：收紧逻辑词表 + 阈值改判规则 + 判别词"""
import pandas as pd
from collections import Counter, defaultdict

CLASSES = ['逻辑思辨', '实证主义', '技术驱动', '人文解释']

DICT = {
 '技术驱动': ['大语言模型', '扩散模型', 'Transformer', 'BERT', '神经网络', '深度学习',
            '机器学习', '知识图谱', '自然语言', '智能体', '推荐系统', '虚假评论', '对话系统',
            '语音识别', '图像识别', '算法设计', '系统设计', '平台开发', '技术实现',
            '分类模型', '预测模型', '生成模型', '模型训练', '数据挖掘', '仿真', '智能客服',
            '语义分析', '情感计算', '计算机视觉', '检索增强', '预训练', '微调', '提示词',
            '多模态', '注意力机制', '个性化推荐', '用户画像', '点击率', '转化率',
            '算法模型', '数据模型', '分析模型', '平台建设', '系统开发', '智能系统',
            '技术开发', '生成算法', '深度学习模型', '算法开发', '系统构建',
            '技术方案', '技术框架', '开发实现'],
 '实证主义': ['问卷', '量表', '实证研究', '实证分析', '实证检验', '实证发现', '实证',
            '回归', '回归分析', '计量', '面板数据', '信度', '效度', '因子分析', '描述性统计',
            '显著性', '假设检验', '结构方程', '实验', '眼动', '行为实验', '对照实验', 'A/B',
            '专利', '统计', '测度', '数据集', '样本', '检验', '方差', '相关系数',
            '使用意愿', '接受模型', 'TAM', 'UTAUT', '空间分布', '区位', '对话测试',
            '双刃剑', '消费者评价', '公平性损失', '量化', '数据统计', '计量分析',
            '博弈', '演化博弈', '数值模拟', '调查', '数据收集', '结构方程模型',
            '评估', '数据驱动分析', '效应', '中介', '调节', '影响因素', '机制检验',
            '意愿', '忠诚', '购买行为', '购买意愿', '重购', '顾客', '上市公司', '年报',
            '经验证据', '实验组', '对照组', '试点', '政策效应', '双重差分', 'DID',
            '工具变量', '稳健性', '内生性', '安慰剂', '机制分析', '异质性', '交互项'],
 '人文解释': ['批判', '意识形态', '规训', '凝视', '伦理', '道德', '价值论', '意义建构',
            '话语', '现象学', '民族志', '参与式观察', '深度访谈', '访谈', '诠释', '异化',
            '权力', '反思', '哲学', '政治经济学', '传统', '生死', '身体', '精神控制',
            '案例研究', '质性', '叙事', '语境', '符号', '文化解读', '批判性', '思辨性反思',
            '可供性', '伦理挑战', '人文', '社会学', '文化分析', '批判路径', '数智悖论',
            '生命', '主体性', '客体化', '数字化生存', '亲密关系', '身份', '算法凝视',
            '隐私危机', '批判性反思', '价值危机', '人文关怀', '道德困境'],
 '逻辑思辨': ['综述', '展望', '探析', '论析', '研究热点', '演进', '变迁', '图景', '述评',
            '评述', '元分析', '知识结构', '发展态势', '研究现状', '对策建议', '体系构建',
            '框架构建', '路径探析', '机理分析', '路径', '机理', '策略', '挑战', '机遇',
            '趋势', '演进', '框架', '体系', '建构', '对策', '赋能', '转型', '变革',
            '审视', '启示', '建议', '革命', '治理', '规制', '监管', '立法', '政策',
            '时代', '演化', '困境', '纾解', '风险', '比较', '实践', '报告', '年度',
            '热点', '视域', '背景', '新内容革命', '发展报告', '探析', '论纲', '刍议',
            '浅析', '思考', '逻辑', '范式', '转向', '重构', '重塑', '再思考'],
}

def score_text(text, title):
    s = str(text); t = str(title)
    sc = {}
    for c in CLASSES:
        sc[c] = sum(s.count(w) + 2 * t.count(w) for w in DICT[c])
    return sc

def assign(sc):
    nonlogic = {c: sc[c] for c in ['实证主义', '技术驱动', '人文解释']}
    mx_nl = max(nonlogic.values())
    if mx_nl <= sc['逻辑思辨']:
        return '逻辑思辨'
    cands = [c for c in ['技术驱动', '实证主义', '人文解释'] if sc[c] == mx_nl]
    if len(cands) == 1:
        return cands[0]
    # 并列：实证优先于技术/人文（实证信号更强）
    for c in ['实证主义', '技术驱动', '人文解释']:
        if c in cands:
            return c
    return '逻辑思辨'

def kappa(y1, y2):
    n = len(y1)
    p_o = sum(a == b for a, b in zip(y1, y2)) / n
    cats = set(y1) | set(y2)
    p_e = 0
    for c in cats:
        p_e += (sum(1 for x in y1 if x == c) / n) * (sum(1 for x in y2 if x == c) / n)
    return (p_o - p_e) / (1 - p_e) if p_e != 1 else 1

cr = pd.read_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_clean_records.csv', encoding='utf-8-sig')
mp = pd.read_csv(r'D:\论文\电商综述\AIGC文创电商\查新文本\_dict_codes_for_compare.csv', encoding='utf-8-sig')
import openpyxl
wb = openpyxl.load_workbook(r'D:\论文\电商综述\AIGC文创电商\查新文本\人工复核表_字典编码_150篇.xlsx')
ws = wb['人工复核表']
manual = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=5).value for r in range(3, 153)}
mp['manual'] = mp['序号'].map(manual)
tr = cr.merge(mp[['题名', 'manual']], on='题名', how='inner')

preds = []
for _, r in tr.iterrows():
    preds.append(assign(score_text(r['摘要'], r['题名'])))
acc = sum(a == b for a, b in zip(preds, tr['manual'])) / len(tr)
kap = kappa(preds, tr['manual'])
print('150篇: 准确率=%.3f kappa=%.3f' % (acc, kap))
cm = defaultdict(Counter)
for t, p in zip(tr['manual'], preds):
    cm[t][p] += 1
print('混淆矩阵(行=人工 列=词典):')
print('%8s' % '', *['%8s' % c for c in CLASSES])
for c in CLASSES:
    print('%8s' % c, *['%8d' % cm[c][x] for x in CLASSES])
for c in CLASSES:
    tp = cm[c][c]
    prec = tp / sum(cm[x][c] for x in CLASSES) if sum(cm[x][c] for x in CLASSES) else 0
    rec = tp / sum(cm[c][x] for x in CLASSES) if sum(cm[c][x] for x in CLASSES) else 0
    f1 = 2 * prec * rec / (prec + rec) if prec + rec else 0
    print('%s: P=%.3f R=%.3f F1=%.3f (tp=%d)' % (c, prec, rec, f1, tp))

allp = []
for _, r in cr.iterrows():
    allp.append(assign(score_text(r['摘要'], r['题名'])))
pc = Counter(allp)
print()
print('全样本分布:', dict(pc))
for c in CLASSES:
    print('  %s: %d (%.1f%%)' % (c, pc[c], 100 * pc[c] / len(allp)))
cr2 = cr.copy()
cr2['范式_校准'] = allp
cr2[['题名', '期刊', '年份', '范式_校准', '摘要']].to_csv(
    r'D:\论文\电商综述\AIGC文创电商\查新文本\_calibrated_codes.csv', index=False, encoding='utf-8-sig')
with open(r'C:\Temp_rend\_cv_result.txt', 'w', encoding='utf-8') as f:
    f.write('词级词典v2 150篇 acc=%.3f kappa=%.3f\n' % (acc, kap))
    f.write('全样本分布: %s\n' % dict(pc))
print('已存档')
